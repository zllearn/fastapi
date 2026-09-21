#!/usr/bin/env python3
"""Build the website's unified static-data database from one consolidated XLSX.

The workbook is the canonical patent-family source.  Each row represents one
IncoPat simple patent family and already contains reviewed owner geography,
owner type, technology, route and three-level industry-chain labels.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = Path(
    "/mnt/d/工作/核聚变卡脖子/IncoPat/incopat核聚变/"
    "IncoPat筛选维度后_已回填企业类型及国家及技术标签.xlsx"
)
DEFAULT_OUTPUT = ROOT / "output" / "统一专利族数据.sqlite3"

FAMILY_SOURCE_FIELDS = (
    "序号", "家族ID", "家族代表公开（公告）号", "完整简单同族成员数",
    "家族国家/地区数量", "家族IPC数量", "家族申请人数", "家族发明人数",
    "家族是否有效", "标题 (中文)", "标题 (英文)", "摘要 (中文)",
    "摘要 (英文)", "首项权利要求-中文", "独立权利要求", "技术功效句",
    "用途", "IPC", "IPC主分类-小类", "IPC主分类-小类(释义)",
    "国民经济分类", "国民经济行业(主)", "新兴产业分类", "新兴产业(主)",
    "申请人", "标准化申请人", "当前权利人", "标准化当前权利人",
    "申请人终属母公司(中文)", "申请人终属母公司(英文)", "申请人类型",
    "申请人国家/地区", "申请人省市代码", "中国申请人地市", "中国申请人区县",
    "当前专利权人地址", "家族引证", "家族被引证", "简单同族",
    "同族国家/地区", "优先权日", "最早优先权日", "优先权国别",
    "首次公开日", "合享价值度", "技术稳定性", "技术先进性", "保护范围",
    "DWPI标题", "DWPI用途", "DWPI优势", "DWPI新颖性", "DWPI详细描述",
    "DWPI技术要点", "DWPI分类号",
)

TECH_SOURCE_FIELDS = (
    "技术特征", "技术分类代码", "技术分类", "具体技术路线代码",
    "具体技术路线", "产业链三级代码", "产业链一级标签",
    "产业链二级标签", "产业链三级标签", "分类依据",
)

OWNER_SOURCE_FIELDS = (
    "当前权利人企业类型（回填）", "当前权利人国家/地区（回填）",
    "当前权利人省/州（仅中国）", "当前权利人城市（仅中国）",
)

HISTORICAL_COUNTRY_MAP = {
    "前苏联": "俄罗斯", "苏联": "俄罗斯", "东德": "德国", "联邦德国": "德国",
    "捷克斯洛伐克": "捷克", "南斯拉夫": "塞尔维亚", "荷属安的列斯": "库拉索",
    "荷属安的列斯（库拉索）": "库拉索", "英属维尔京群岛": "维尔京群岛",
}

CHINA_REGIONS = {
    "北京", "天津", "上海", "重庆", "河北", "山西", "辽宁", "吉林", "黑龙江",
    "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南",
    "广东", "海南", "四川", "贵州", "云南", "陕西", "甘肃", "青海", "台湾",
    "内蒙古", "广西", "西藏", "宁夏", "新疆", "香港", "澳门",
}

VALID_TYPES = {"私企", "国企", "高校", "研究所", "政府", "个人", "待核验", "主体信息缺失"}

def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalized(value: object) -> str:
    value = unicodedata.normalize("NFKC", text(value)).casefold()
    return re.sub(r"[\s\-—_·•,，.。;；:：()（）\[\]【】]+", "", value)


def split_values(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[;；\r\n]+", text(value)) if part.strip()]


def split_geo_values(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[;,，；\r\n]+", text(value)) if part.strip()]


def clean_region(value: object) -> str:
    result = text(value)
    while len(result) > 1 and ((result[0], result[-1]) in {("[", "]"), ("【", "】")}):
        result = result[1:-1].strip()
    result = result.removesuffix("省").removesuffix("市")
    aliases = {
        "内蒙古自治区": "内蒙古", "广西壮族自治区": "广西", "西藏自治区": "西藏",
        "宁夏回族自治区": "宁夏", "新疆维吾尔自治区": "新疆",
        "香港特别行政区": "香港", "澳门特别行政区": "澳门",
        "中国台湾": "台湾", "中国香港": "香港", "中国澳门": "澳门",
    }
    return aliases.get(result, result)


def normalize_country(value: object) -> tuple[str, str, str]:
    raw = text(value)
    if raw in {"台湾", "中国台湾"}:
        return raw, "中国", "台湾"
    if raw in {"香港", "中国香港"}:
        return raw, "中国", "香港"
    if raw in {"澳门", "中国澳门"}:
        return raw, "中国", "澳门"
    if not raw or raw in {"无法判定", "未识别", "国家未识别", "ZZ"}:
        return raw, "国家未识别", ""
    return raw, HISTORICAL_COUNTRY_MAP.get(raw, raw), ""


def extract_year(*values: object) -> int | None:
    for value in values:
        match = re.search(r"(?:18|19|20)\d{2}", text(value))
        if match:
            return int(match.group())
    return None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def aligned(values: list[str], length: int, default: str) -> list[str]:
    if length <= 0:
        return []
    if len(values) == length:
        return values
    if len(values) == 1:
        return values * length
    return [values[index] if index < len(values) else default for index in range(length)]


def choose_mode(counter: Counter[str], default: str) -> str:
    if not counter:
        return default
    return sorted(counter.items(), key=lambda item: (-item[1], item[0]))[0][0]


def create_schema(connection: sqlite3.Connection) -> None:
    family_columns = ",\n".join(f'"{field}" TEXT' for field in FAMILY_SOURCE_FIELDS)
    connection.executescript(f'''
        PRAGMA journal_mode=OFF;
        PRAGMA synchronous=OFF;
        PRAGMA temp_store=MEMORY;
        PRAGMA foreign_keys=ON;
        CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL) WITHOUT ROWID;
        CREATE TABLE families(
            family_id TEXT PRIMARY KEY,
            source_patent_id INTEGER NOT NULL UNIQUE,
            priority_year INTEGER,
            {family_columns}
        );
        CREATE TABLE family_tech(
            family_id TEXT PRIMARY KEY REFERENCES families(family_id),
            processing_status TEXT NOT NULL,
            technical_feature TEXT NOT NULL,
            technology_code TEXT NOT NULL,
            technology_label TEXT NOT NULL,
            route_code TEXT NOT NULL,
            route_label TEXT NOT NULL,
            chain_level3_code TEXT NOT NULL,
            chain_level1 TEXT NOT NULL,
            chain_level2 TEXT NOT NULL,
            chain_level3 TEXT NOT NULL,
            rationale TEXT NOT NULL,
            confidence TEXT NOT NULL,
            error_message TEXT
        );
        CREATE TABLE entities(
            entity_id INTEGER PRIMARY KEY,
            representative_name TEXT NOT NULL UNIQUE,
            enterprise_type TEXT NOT NULL,
            source_part TEXT,
            source_row INTEGER
        );
        CREATE TABLE family_entities(
            family_id TEXT NOT NULL REFERENCES families(family_id),
            entity_order INTEGER NOT NULL,
            entity_id INTEGER REFERENCES entities(entity_id),
            entity_name TEXT NOT NULL,
            entity_source TEXT NOT NULL,
            enterprise_type TEXT NOT NULL,
            match_status TEXT NOT NULL,
            PRIMARY KEY(family_id, entity_order)
        ) WITHOUT ROWID;
        CREATE TABLE family_applicants(
            family_id TEXT NOT NULL REFERENCES families(family_id),
            applicant_order INTEGER NOT NULL,
            applicant_name TEXT NOT NULL,
            entity_id INTEGER REFERENCES entities(entity_id),
            enterprise_type TEXT NOT NULL,
            match_status TEXT NOT NULL,
            PRIMARY KEY(family_id, applicant_order)
        ) WITHOUT ROWID;
        CREATE TABLE owner_name_locations(
            owner_name TEXT PRIMARY KEY,
            entity_id INTEGER REFERENCES entities(entity_id),
            raw_country TEXT,
            country TEXT NOT NULL,
            province TEXT,
            city TEXT,
            location_source TEXT NOT NULL,
            confidence TEXT,
            evidence_count INTEGER NOT NULL
        ) WITHOUT ROWID;
        CREATE TABLE entity_locations(
            entity_id INTEGER PRIMARY KEY REFERENCES entities(entity_id),
            raw_country TEXT,
            country TEXT NOT NULL,
            province TEXT,
            city TEXT,
            location_source TEXT NOT NULL,
            evidence_count INTEGER NOT NULL
        ) WITHOUT ROWID;
        CREATE TABLE family_applicant_geographies(
            family_id TEXT NOT NULL REFERENCES families(family_id),
            geo_order INTEGER NOT NULL,
            raw_country TEXT,
            country TEXT NOT NULL,
            province TEXT,
            city TEXT,
            district TEXT,
            PRIMARY KEY(family_id, geo_order)
        ) WITHOUT ROWID;
        CREATE TABLE family_entity_comparison(
            family_id TEXT PRIMARY KEY REFERENCES families(family_id),
            applicant_entities TEXT,
            current_owner_entities TEXT,
            both_present INTEGER NOT NULL,
            raw_sets_equal INTEGER NOT NULL
        ) WITHOUT ROWID;
        CREATE TABLE quality_metrics(
            metric TEXT PRIMARY KEY,
            value TEXT NOT NULL,
            status TEXT NOT NULL,
            note TEXT
        ) WITHOUT ROWID;
    ''')


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    source = args.input.resolve()
    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".building")
    if temporary.exists():
        temporary.unlink()

    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["简单专利族合并"]
    headers = [text(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    required = set(FAMILY_SOURCE_FIELDS) | set(TECH_SOURCE_FIELDS) | set(OWNER_SOURCE_FIELDS)
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"输入表缺少字段: {missing}")
    position = {name: headers.index(name) for name in required}

    connection = sqlite3.connect(temporary)
    create_schema(connection)
    family_sql = (
        f'INSERT INTO families VALUES ({",".join("?" for _ in range(3 + len(FAMILY_SOURCE_FIELDS)))})'
    )
    tech_sql = f'INSERT INTO family_tech VALUES ({",".join("?" for _ in range(14))})'

    families_batch: list[tuple] = []
    tech_batch: list[tuple] = []
    applicant_geo_batch: list[tuple] = []
    comparison_batch: list[tuple] = []
    pending_entities: list[dict] = []
    pending_applicants: list[tuple[str, list[str]]] = []
    family_ids: set[str] = set()
    metrics = Counter()

    # entity_key -> evidence. Standardized owner is used only as a grouping key;
    # the most frequent raw current-owner name remains the display name.
    entity_names: dict[str, Counter[str]] = defaultdict(Counter)
    entity_types: dict[str, Counter[str]] = defaultdict(Counter)
    entity_locations: dict[str, Counter[tuple[str, str, str, str]]] = defaultdict(Counter)
    owner_locations: dict[str, Counter[tuple[str, str, str, str, str]]] = defaultdict(Counter)

    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        family_id = text(row[position["家族ID"]])
        if not family_id or family_id in family_ids:
            raise ValueError(f"家族ID为空或重复，Excel行 {source_row}: {family_id!r}")
        family_ids.add(family_id)
        priority_year = extract_year(
            row[position["最早优先权日"]], row[position["优先权日"]], row[position["首次公开日"]]
        )
        family_values = tuple(text(row[position[field]]) for field in FAMILY_SOURCE_FIELDS)
        families_batch.append((family_id, source_row - 1, priority_year, *family_values))
        technology_code = text(row[position["技术分类代码"]])
        technology_label = text(row[position["技术分类"]])
        if technology_code == "B7" and technology_label == "用聚变支撑技术":
            technology_label = "通用聚变支撑技术"
        tech_batch.append((
            family_id, "complete", text(row[position["技术特征"]]),
            technology_code, technology_label,
            text(row[position["具体技术路线代码"]]), text(row[position["具体技术路线"]]),
            text(row[position["产业链三级代码"]]), text(row[position["产业链一级标签"]]),
            text(row[position["产业链二级标签"]]), text(row[position["产业链三级标签"]]),
            text(row[position["分类依据"]]), "已复核", "",
        ))

        # Applicant geography keeps the applicant-based benchmark convention.
        countries = split_geo_values(row[position["申请人国家/地区"]]) or ["国家未识别"]
        provinces = [
            region for value in split_geo_values(row[position["申请人省市代码"]])
            if (region := clean_region(value)) in CHINA_REGIONS
        ]
        cities = [clean_region(value) for value in split_geo_values(row[position["中国申请人地市"]])]
        districts = [clean_region(value) for value in split_geo_values(row[position["中国申请人区县"]])]
        china_index = 0
        seen_geo: set[tuple[str, str, str, str]] = set()
        for raw_country in countries:
            original, country, implied_province = normalize_country(raw_country)
            province = city = district = ""
            if country == "中国":
                province = implied_province or (provinces[china_index] if china_index < len(provinces) else "")
                city = cities[china_index] if china_index < len(cities) else ""
                district = districts[china_index] if china_index < len(districts) else ""
                if province in {"北京", "天津", "上海", "重庆", "香港", "澳门"} and not city:
                    city = province
                china_index += 1
            key = (country, province, city, district)
            if key in seen_geo:
                continue
            seen_geo.add(key)
            applicant_geo_batch.append((family_id, len(seen_geo), original, country, province, city, district))

        raw_owners = split_values(row[position["当前权利人"]])
        standardized = split_values(row[position["标准化当前权利人"]])
        owner_types = aligned(split_values(row[position["当前权利人企业类型（回填）"]]), len(raw_owners), "待核验")
        owner_countries = aligned(split_values(row[position["当前权利人国家/地区（回填）"]]), len(raw_owners), "国家未识别")
        owner_provinces = aligned(split_values(row[position["当前权利人省/州（仅中国）"]]), len(raw_owners), "不适用")
        owner_cities = aligned(split_values(row[position["当前权利人城市（仅中国）"]]), len(raw_owners), "不适用")
        if raw_owners and len(standardized) != len(raw_owners):
            metrics["标准化当前权利人与当前权利人数量不一致"] += 1
            standardized = raw_owners
        elif not standardized:
            standardized = raw_owners

        entity_source = "当前权利人"
        if not raw_owners:
            fallbacks = (
                ("申请人终属母公司(中文)", row[position["申请人终属母公司(中文)"]]),
                ("申请人终属母公司(英文)", row[position["申请人终属母公司(英文)"]]),
                ("标准化申请人", row[position["标准化申请人"]]),
            )
            entity_source, chosen = next(((name, value) for name, value in fallbacks if text(value)), ("无可用主体字段", "主体未识别"))
            raw_owners = split_values(chosen) or ["主体未识别"]
            standardized = raw_owners
            owner_types = aligned(split_values(row[position["当前权利人企业类型（回填）"]]), len(raw_owners), "主体信息缺失")
            owner_countries = aligned(split_values(row[position["当前权利人国家/地区（回填）"]]), len(raw_owners), "国家未识别")
            owner_provinces = aligned(split_values(row[position["当前权利人省/州（仅中国）"]]), len(raw_owners), "不适用")
            owner_cities = aligned(split_values(row[position["当前权利人城市（仅中国）"]]), len(raw_owners), "不适用")
            metrics["无当前权利人家族"] += 1

        family_entity_items = []
        for index, owner in enumerate(raw_owners):
            canonical = standardized[index] if index < len(standardized) and standardized[index] else owner
            entity_key = normalized(canonical) or normalized(owner) or f"missing:{family_id}:{index}"
            enterprise_type = owner_types[index] if index < len(owner_types) else "待核验"
            if enterprise_type not in VALID_TYPES:
                enterprise_type = "待核验"
            raw_country = owner_countries[index] if index < len(owner_countries) else "国家未识别"
            original_country, country, implied_province = normalize_country(raw_country)
            province = clean_region(owner_provinces[index] if index < len(owner_provinces) else "")
            city = clean_region(owner_cities[index] if index < len(owner_cities) else "")
            if province in {"不适用", "国家未识别"}: province = ""
            if city in {"不适用", "国家未识别"}: city = ""
            if country == "中国":
                province = implied_province or province
                if province in {"北京", "天津", "上海", "重庆", "香港", "澳门"} and not city:
                    city = province
            else:
                province = city = ""
            entity_names[entity_key][owner] += 1
            entity_types[entity_key][enterprise_type] += 1
            entity_locations[entity_key][(original_country, country, province, city)] += 1
            if entity_source == "当前权利人" and owner != "主体未识别":
                owner_locations[owner][(original_country, country, province, city, entity_key)] += 1
            family_entity_items.append((owner, entity_key, entity_source, enterprise_type))
        pending_entities.append({"family_id": family_id, "items": family_entity_items})

        applicants = split_values(row[position["标准化申请人"]]) or split_values(row[position["申请人"]])
        pending_applicants.append((family_id, applicants))
        applicant_set = {normalized(value) for value in applicants if normalized(value)}
        owner_set = {normalized(value) for value in raw_owners if value != "主体未识别" and normalized(value)}
        comparison_batch.append((
            family_id, "; ".join(sorted(applicant_set)), "; ".join(sorted(owner_set)),
            int(bool(applicant_set and owner_set)), int(bool(applicant_set and owner_set and applicant_set == owner_set)),
        ))

        if len(families_batch) >= 1000:
            connection.executemany(family_sql, families_batch)
            connection.executemany(tech_sql, tech_batch)
            connection.executemany("INSERT INTO family_applicant_geographies VALUES (?,?,?,?,?,?,?)", applicant_geo_batch)
            connection.executemany("INSERT INTO family_entity_comparison VALUES (?,?,?,?,?)", comparison_batch)
            families_batch.clear(); tech_batch.clear(); applicant_geo_batch.clear(); comparison_batch.clear()

    if families_batch:
        connection.executemany(family_sql, families_batch)
        connection.executemany(tech_sql, tech_batch)
        connection.executemany("INSERT INTO family_applicant_geographies VALUES (?,?,?,?,?,?,?)", applicant_geo_batch)
        connection.executemany("INSERT INTO family_entity_comparison VALUES (?,?,?,?,?)", comparison_batch)

    # Stable entity IDs are assigned by normalized grouping key.
    entity_ids: dict[str, int] = {}
    entity_rows = []
    location_rows = []
    representative_to_key: dict[str, str] = {}
    for entity_id, key in enumerate(sorted(entity_names), start=1):
        representative = choose_mode(entity_names[key], key)
        # Extremely rare standardized collisions must not violate the display-name unique key.
        if representative in representative_to_key and representative_to_key[representative] != key:
            representative = f"{representative}（{entity_id}）"
        representative_to_key[representative] = key
        enterprise_type = choose_mode(entity_types[key], "待核验")
        raw_country, country, province, city = choose_mode(entity_locations[key], ("", "国家未识别", "", ""))
        entity_ids[key] = entity_id
        entity_rows.append((entity_id, representative, enterprise_type, "合并Excel标准化当前权利人归并", None))
        location_rows.append((entity_id, raw_country, country, province, city, "Excel回填字段", sum(entity_locations[key].values())))
    connection.executemany("INSERT INTO entities VALUES (?,?,?,?,?)", entity_rows)
    connection.executemany("INSERT INTO entity_locations VALUES (?,?,?,?,?,?,?)", location_rows)

    family_entity_rows = []
    for item in pending_entities:
        for order, (owner, key, source_name, enterprise_type) in enumerate(item["items"], start=1):
            family_entity_rows.append((
                item["family_id"], order, entity_ids[key], owner, source_name, enterprise_type,
                "Excel字段直接归并" if source_name == "当前权利人" else "Excel回退链",
            ))
    connection.executemany("INSERT INTO family_entities VALUES (?,?,?,?,?,?,?)", family_entity_rows)

    owner_location_rows = []
    for owner, evidence in sorted(owner_locations.items()):
        raw_country, country, province, city, key = choose_mode(evidence, ("", "国家未识别", "", "", ""))
        owner_location_rows.append((
            owner, entity_ids.get(key), raw_country, country, province, city,
            "Excel回填字段", "已回填", sum(evidence.values()),
        ))
    connection.executemany("INSERT INTO owner_name_locations VALUES (?,?,?,?,?,?,?,?,?)", owner_location_rows)

    applicant_rows = []
    key_by_normalized_display = {normalized(row[1]): row[0] for row in entity_rows}
    type_by_id = {row[0]: row[2] for row in entity_rows}
    for family_id, applicants in pending_applicants:
        for order, applicant in enumerate(applicants, start=1):
            entity_id = key_by_normalized_display.get(normalized(applicant))
            applicant_rows.append((
                family_id, order, entity_id and applicant or applicant, entity_id,
                type_by_id.get(entity_id, "待核验"), "当前权利人实体同名匹配" if entity_id else "未匹配",
            ))
    connection.executemany("INSERT INTO family_applicants VALUES (?,?,?,?,?,?)", applicant_rows)

    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    metadata = {
        "dataset_title": "IncoPat核聚变统一专利族数据",
        "data_source": "IncoPat consolidated XLSX",
        "family_scope": "简单专利族",
        "relevance_scope": "全部家族纳入；前端按既定规则展示B1-B9",
        "enterprise_grouping": "当前权利人展示；标准化当前权利人仅用于主体归并；缺失时采用回退链",
        "geography_scope": "国家比较采用申请人地理；企业画像采用当前权利人回填地理",
        "historical_country_rule": json.dumps(HISTORICAL_COUNTRY_MAP, ensure_ascii=False, sort_keys=True),
        "generated_at_utc": generated,
        "source_xlsx": str(source),
        "source_xlsx_sha256": sha256(source),
    }
    connection.executemany("INSERT INTO metadata VALUES (?,?)", metadata.items())

    quality = [
        ("简单专利族数", str(len(family_ids)), "通过" if len(family_ids) == 68151 else "异常", "家族ID唯一"),
        ("无当前权利人家族数", str(metrics["无当前权利人家族"]), "信息", "按终属母公司/申请人/主体未识别回退"),
        ("标准化权利人数量不一致家族数", str(metrics["标准化当前权利人与当前权利人数量不一致"]), "信息", "不一致时按当前权利人独立归并"),
        ("权利人实体数", str(len(entity_rows)), "信息", "标准化字段后台归并，页面展示当前权利人"),
        ("当前权利人名称数", str(len(owner_location_rows)), "信息", "含地理回填"),
    ]
    connection.executemany("INSERT INTO quality_metrics VALUES (?,?,?,?)", quality)
    connection.executescript('''
        CREATE INDEX idx_family_priority_year ON families(priority_year);
        CREATE INDEX idx_tech_code ON family_tech(technology_code);
        CREATE INDEX idx_tech_chain1 ON family_tech(chain_level1);
        CREATE INDEX idx_family_entity_entity ON family_entities(entity_id, family_id);
        CREATE INDEX idx_family_entity_type ON family_entities(enterprise_type, family_id);
        CREATE INDEX idx_family_applicant_entity ON family_applicants(entity_id, family_id);
        CREATE INDEX idx_family_applicant_type ON family_applicants(enterprise_type, family_id);
        CREATE INDEX idx_owner_location_country ON owner_name_locations(country);
        CREATE INDEX idx_entity_location_country ON entity_locations(country);
        CREATE INDEX idx_applicant_geo_country ON family_applicant_geographies(country, family_id);
        CREATE INDEX idx_applicant_geo_china ON family_applicant_geographies(province, city, family_id);
        CREATE VIEW v_family_complete AS
            SELECT f.*, t.technical_feature, t.technology_code, t.technology_label,
                   t.route_code, t.route_label, t.chain_level3_code,
                   t.chain_level1, t.chain_level2, t.chain_level3,
                   t.rationale AS classification_rationale,
                   t.confidence AS classification_confidence
            FROM families f JOIN family_tech t USING(family_id);
        CREATE VIEW v_family_current_owner_geography AS
            SELECT fe.family_id, fe.entity_order, fe.entity_id, fe.entity_name,
                   fe.enterprise_type, el.raw_country, el.country, el.province,
                   el.city, el.location_source
            FROM family_entities fe LEFT JOIN entity_locations el USING(entity_id);
        PRAGMA optimize;
    ''')
    connection.commit()

    checks = {
        "families": connection.execute("SELECT COUNT(*) FROM families").fetchone()[0],
        "tech": connection.execute("SELECT COUNT(*) FROM family_tech").fetchone()[0],
        "entities": connection.execute("SELECT COUNT(*) FROM entities").fetchone()[0],
        "current_owner_families": connection.execute(
            "SELECT COUNT(DISTINCT family_id) FROM family_entities WHERE entity_source='当前权利人'"
        ).fetchone()[0],
        "integrity": connection.execute("PRAGMA integrity_check").fetchone()[0],
    }
    connection.close()
    workbook.close()
    if checks["families"] != 68151 or checks["tech"] != checks["families"] or checks["integrity"] != "ok":
        raise RuntimeError(f"构建校验失败: {checks}")
    os.replace(temporary, output)
    print(json.dumps({"output": str(output), "source": str(source), **checks}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
