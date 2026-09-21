#!/usr/bin/env python3
"""Build the static topic -> patent -> current owner -> enterprise index."""

from __future__ import annotations

import argparse
import json
import sqlite3
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PIPELINE_ROOT = Path(__file__).resolve().parent
PROJECTS_ROOT = PIPELINE_ROOT.parent
DEFAULT_WORKBOOK = (
    PROJECTS_ROOT
    / "analytics/incopat_run/incopat_simple_families_topic_indicators.xlsx"
)
DEFAULT_DATABASE = PIPELINE_ROOT / "output/unified_patent_families.sqlite3"
DEFAULT_OUTPUT = PIPELINE_ROOT / "output/site-payloads/topic-enterprise-index.js"


def clean(value: object) -> str:
    return str(value or "").strip()


def load_topic_patents(workbook_path: Path) -> tuple[dict[str, dict], dict[int, list[str]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook["02_doc_topics"]
    rows = worksheet.iter_rows(values_only=True)
    headers = {name: index for index, name in enumerate(next(rows))}
    required = {
        "source_family_id",
        "source_display_key",
        "source_publication_date",
        "topic_id",
        "TI_raw",
    }
    missing = required.difference(headers)
    if missing:
        workbook.close()
        raise ValueError(f"02_doc_topics 缺少字段: {', '.join(sorted(missing))}")

    patents: dict[str, dict] = {}
    topic_families: dict[int, list[str]] = defaultdict(list)
    for row in rows:
        family_id = clean(row[headers["source_family_id"]])
        if not family_id:
            continue
        topic_id = int(row[headers["topic_id"]])
        patents[family_id] = {
            "familyId": family_id,
            "publication": clean(row[headers["source_display_key"]]) or family_id,
            "date": clean(row[headers["source_publication_date"]]),
            "title": clean(row[headers["TI_raw"]]) or "标题未收录",
        }
        topic_families[topic_id].append(family_id)
    workbook.close()
    return patents, topic_families


def load_current_owners(
    database_path: Path, family_ids: list[str]
) -> dict[str, list[dict]]:
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    connection.execute("CREATE TEMP TABLE wanted_family (family_id TEXT PRIMARY KEY)")
    connection.executemany(
        "INSERT OR IGNORE INTO wanted_family VALUES (?)",
        ((family_id,) for family_id in family_ids),
    )
    rows = connection.execute(
        """
        SELECT fe.family_id,
               fe.entity_order,
               fe.entity_id,
               fe.entity_name AS current_owner,
               fe.enterprise_type AS owner_type,
               fe.match_status,
               e.representative_name AS company_name,
               e.enterprise_type AS company_type,
               l.country,
               l.province,
               l.city
        FROM family_entities AS fe
        JOIN wanted_family AS wanted USING (family_id)
        LEFT JOIN entities AS e USING (entity_id)
        LEFT JOIN entity_locations AS l USING (entity_id)
        WHERE fe.entity_source = '当前权利人'
        ORDER BY fe.family_id, fe.entity_order
        """
    )
    owners_by_family: dict[str, list[dict]] = defaultdict(list)
    seen: dict[str, set[tuple[object, str]]] = defaultdict(set)
    for row in rows:
        family_id = clean(row["family_id"])
        current_owner = clean(row["current_owner"])
        identity = (row["entity_id"], current_owner)
        if not current_owner or identity in seen[family_id]:
            continue
        seen[family_id].add(identity)
        company_name = clean(row["company_name"])
        company_type = clean(row["company_type"] or row["owner_type"] or "待核验")
        country = clean(row["country"])
        can_open_profile = bool(
            row["entity_id"]
            and company_name
            and country == "中国"
        )
        owners_by_family[family_id].append(
            {
                "entityId": row["entity_id"],
                "currentOwner": current_owner,
                "type": company_type,
                "matchStatus": clean(row["match_status"]),
                "company": company_name if can_open_profile else "",
                "country": country,
                "province": clean(row["province"]),
                "city": clean(row["city"]),
            }
        )
    connection.close()
    return owners_by_family


def load_chain_dimensions(database_path: Path, family_ids: list[str]) -> dict[str, dict[str, str]]:
    connection = sqlite3.connect(database_path)
    connection.execute("CREATE TEMP TABLE wanted_chain_family (family_id TEXT PRIMARY KEY)")
    connection.executemany(
        "INSERT OR IGNORE INTO wanted_chain_family VALUES (?)",
        ((family_id,) for family_id in family_ids),
    )
    rows = connection.execute(
        """
        SELECT ft.family_id, ft.chain_level1, ft.chain_level2
        FROM family_tech AS ft
        JOIN wanted_chain_family AS wanted USING (family_id)
        """
    )
    chain_dimensions = {
        clean(family_id): {
            "chainLevel1": clean(chain_level1),
            "chainLevel2": clean(chain_level2),
        }
        for family_id, chain_level1, chain_level2 in rows
    }
    connection.close()
    return chain_dimensions


def build_payload(
    workbook_path: Path, database_path: Path
) -> dict:
    patents, topic_families = load_topic_patents(workbook_path)
    owners_by_family = load_current_owners(database_path, list(patents))
    chain_dimensions = load_chain_dimensions(database_path, list(patents))
    topics: dict[str, dict] = {}
    all_owner_families = 0
    all_mapped_families = 0
    profile_names: set[str] = set()

    for topic_id, family_ids in sorted(topic_families.items()):
        topic_patents: list[dict] = []
        company_aggregate: dict[str, dict] = {}
        with_current_owner = 0
        mapped_families = 0

        for family_id in family_ids:
            owner_rows = owners_by_family.get(family_id, [])
            if owner_rows:
                with_current_owner += 1
            mapped_company_names: set[str] = set()
            for owner in owner_rows:
                company_name = owner["company"]
                if not company_name:
                    continue
                mapped_company_names.add(company_name)
                aggregate = company_aggregate.setdefault(
                    company_name,
                    {
                        "name": company_name,
                        "type": owner["type"],
                        "country": owner["country"],
                        "province": owner["province"],
                        "city": owner["city"],
                        "aliases": set(),
                        "patentIds": set(),
                    },
                )
                aggregate["aliases"].add(owner["currentOwner"])
                aggregate["patentIds"].add(family_id)
                profile_names.add(company_name)
            if mapped_company_names:
                mapped_families += 1
            patent = {
                **patents[family_id],
                **chain_dimensions.get(
                    family_id,
                    {"chainLevel1": "", "chainLevel2": ""},
                ),
                "owners": [
                    {
                        "currentOwner": owner["currentOwner"],
                        "type": owner["type"],
                        "company": owner["company"],
                    }
                    for owner in owner_rows
                ],
            }
            topic_patents.append(patent)

        companies = []
        for aggregate in company_aggregate.values():
            patent_ids = sorted(aggregate.pop("patentIds"))
            aliases = sorted(aggregate.pop("aliases"))
            companies.append(
                {
                    **aggregate,
                    "aliases": aliases,
                    "patentCount": len(patent_ids),
                    "patentIds": patent_ids,
                }
            )
        companies.sort(key=lambda item: (-item["patentCount"], item["name"]))
        topic_patents.sort(
            key=lambda item: (item["date"], item["familyId"]), reverse=True
        )
        topics[str(topic_id)] = {
            "patentCount": len(family_ids),
            "withCurrentOwner": with_current_owner,
            "mappedFamilies": mapped_families,
            "companyCount": len(companies),
            "companies": companies,
            "patents": topic_patents,
        }
        all_owner_families += with_current_owner
        all_mapped_families += mapped_families

    return {
        "schemaVersion": 3,
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "method": {
            "topicJoinKey": "02_doc_topics.source_family_id = family_entities.family_id",
            "ownerScope": "family_entities.entity_source=当前权利人",
            "enterpriseScope": "country=中国，不限制主体类型",
            "enterpriseName": "entities.representative_name",
            "chainSource": "family_tech.chain_level1 + family_tech.chain_level2",
            "chainScope": "产业链二级指标（不适用不进入河流）",
        },
        "summary": {
            "topicCount": len(topics),
            "patentCount": len(patents),
            "withCurrentOwner": all_owner_families,
            "mappedFamilies": all_mapped_families,
            "enterpriseCount": len(profile_names),
        },
        "topics": topics,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    payload = build_payload(args.workbook, args.database)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    args.output.write_text(
        "window.FRONTIER_ENTERPRISE_INDEX=" + encoded + ";\n",
        encoding="utf-8",
    )
    print(
        f"Built {args.output} ({args.output.stat().st_size:,} bytes, "
        f"{payload['summary']['patentCount']:,} patents, "
        f"{payload['summary']['enterpriseCount']:,} enterprises)"
    )


if __name__ == "__main__":
    main()
