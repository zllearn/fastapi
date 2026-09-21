#!/usr/bin/env python3
"""Build the verified China enterprise directory used by the static demo site."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = ROOT.parent / "data" / "enterprise_directory.xlsx"
DEFAULT_OUTPUT = ROOT / "output" / "site-payloads" / "enterprise-directory-master.js"


def normalize_name(value: str) -> str:
    text = unicodedata.normalize("NFKC", value).lower()
    text = re.sub(r"[\s\[\]【】()（）·,，.。'\"“”‘’\-—_]", "", text)
    return re.sub(r"有限责任公司$", "有限公司", text)


def clean(value) -> str:
    return str(value or "").strip()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("output", nargs="?", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    sheet = workbook["企业总名单"]
    headers = [clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    column = {name: index for index, name in enumerate(headers)}
    required = ["企业中文名称", "国家/地区", "省份", "城市", "数据来源", "企业核验状态", "相关专利族数"]
    missing = [name for name in required if name not in column]
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    merged: dict[str, dict] = {}
    source_rows = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = clean(row[column["企业中文名称"]])
        if not name:
            continue
        source_rows += 1
        key = normalize_name(name)
        candidate = {
            "name": name,
            "aliases": [name],
            "country": clean(row[column["国家/地区"]]) or "中国",
            "province": clean(row[column["省份"]]),
            "city": clean(row[column["城市"]]),
            "source": clean(row[column["数据来源"]]),
            "verificationStatus": clean(row[column["企业核验状态"]]),
            "sourcePatentCount": int(row[column["相关专利族数"]] or 0),
            "ownerScope": clean(row[column["IncoPat主体口径"]]) if "IncoPat主体口径" in column else "",
            "matchedCurrentOwners": clean(row[column["当前权利人匹配名称"]]) if "当前权利人匹配名称" in column else "",
            "currentOwnerPatentCount": int(row[column["当前权利人专利族数"]] or 0) if "当前权利人专利族数" in column else 0,
            "calibrationNote": clean(row[column["校准说明"]]) if "校准说明" in column else "",
        }
        if key not in merged:
            merged[key] = candidate
            continue
        current = merged[key]
        if name not in current["aliases"]:
            current["aliases"].append(name)
        current["sourcePatentCount"] = max(current["sourcePatentCount"], candidate["sourcePatentCount"])
        current["currentOwnerPatentCount"] = max(current["currentOwnerPatentCount"], candidate["currentOwnerPatentCount"])
        for field in ("country", "province", "city", "source", "verificationStatus", "ownerScope", "matchedCurrentOwners", "calibrationNote"):
            if not current[field] and candidate[field]:
                current[field] = candidate[field]

    entries = sorted(merged.values(), key=lambda item: item["name"])
    payload = {
        "schemaVersion": 2,
        "updatedAt": "2026-09-09",
        "sourceFile": args.source.name,
        "sourceRows": source_rows,
        "enterpriseCount": len(entries),
        "deduplication": "NFKC + 全半角括号/空格/常见标点 + 有限责任公司后缀归一",
        "entries": entries,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "window.ENTERPRISE_DIRECTORY_MASTER=" + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "sourceRows": source_rows,
        "enterpriseCount": len(entries),
        "output": str(args.output),
        "bytes": args.output.stat().st_size,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
