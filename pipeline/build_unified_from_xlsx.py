#!/usr/bin/env python3
"""Build the website's unified static-data database from one consolidated XLSX.

The workbook is the canonical patent-family source.  Each row represents one
IncoPat simple patent family and already contains reviewed owner geography,
owner type, technology, route and three-level industry-chain labels.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent))

from backend import config, schema  # noqa: E402
from backend.database import Connection, connect  # noqa: E402

DEFAULT_INPUT = ROOT.parent / "data" / "incopat_patent_families.xlsx"

# Column names are identical to the upstream workbook headers, so the canonical
# list lives with the MySQL DDL and is re-exported here.
FAMILY_SOURCE_FIELDS = schema.FAMILY_SOURCE_FIELDS

TECH_SOURCE_FIELDS = (
    "技术特征", "技术分类代码", "技术分类", "具体技术路线代码",
    "具体技术路线", "产业链三级代码", "产业链一级标签",
    "产业链二级标签", "产业链三级标签", "分类依据",
)

OWNER_SOURCE_FIELDS = (
    "当前权利人企业类型（回填）", "当前权利人国家/地区（回填）",
    "当前权利人省/州（仅中国）", "当前权利人城市（仅中国）",
)

HISTORICAL_COUNTRY_MAP = {
    "前苏联": "俄罗斯", "苏联": "俄罗斯", "东德": "德国", "联邦德国": "德国",
    "捷克斯洛伐克": "捷克", "南斯拉夫": "塞尔维亚", "荷属安的列斯": "库拉索",
    "荷属安的列斯（库拉索）": "库拉索", "英属维尔京群岛": "维尔京群岛",
}

CHINA_REGIONS = {
    "北京", "天津", "上海", "重庆", "河北", "山西", "辽宁", "吉林", "黑龙江",
    "江苏", "浙江", "安徽", "福建", "江西", "山东", "河南", "湖北", "湖南",
    "广东", "海南", "四川", "贵州", "云南", "陕西", "甘肃", "青海", "台湾",
    "内蒙古", "广西", "西藏", "宁夏", "新疆", "香港", "澳门",
}

VALID_TYPES = {"私企", "国企", "高校", "研究所", "政府", "个人", "待核验", "主体信息缺失"}

def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def normalized(value: object) -> str:
    value = unicodedata.normalize("NFKC", text(value)).casefold()
    return re.sub(r"[\s\-—_·•,，.。;；:：()（）\[\]【】]+", "", value)


def split_values(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[;；\r\n]+", text(value)) if part.strip()]


def split_geo_values(value: object) -> list[str]:
    return [part.strip() for part in re.split(r"[;,，；\r\n]+", text(value)) if part.strip()]


def clean_region(value: object) -> str:
    result = text(value)
    while len(result) > 1 and ((result[0], result[-1]) in {("[", "]"), ("【", "】")}):
        result = result[1:-1].strip()
    result = result.removesuffix("省").removesuffix("市")
    aliases = {
        "内蒙古自治区": "内蒙古", "广西壮族自治区": "广西", "西藏自治区": "西藏",
        "宁夏回族自治区": "宁夏", "新疆维吾尔自治区": "新疆",
        "香港特别行政区": "香港", "澳门特别行政区": "澳门",
        "中国台湾": "台湾", "中国香港": "香港", "中国澳门": "澳门",
    }
    return aliases.get(result, result)


def normalize_country(value: object) -> tuple[str, str, str]:
    raw = text(value)
    if raw in {"台湾", "中国台湾"}:
        return raw, "中国", "台湾"
    if raw in {"香港", "中国香港"}:
        return raw, "中国", "香港"
    if raw in {"澳门", "中国澳门"}:
        return raw, "中国", "澳门"
    if not raw or raw in {"无法判定", "未识别", "国家未识别", "ZZ"}:
        return raw, "国家未识别", ""
    return raw, HISTORICAL_COUNTRY_MAP.get(raw, raw), ""


def extract_year(*values: object) -> int | None:
    for value in values:
        match = re.search(r"(?:18|19|20)\d{2}", text(value))
        if match:
            return int(match.group())
    return None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def aligned(values: list[str], length: int, default: str) -> list[str]:
    if length <= 0:
        return []
    if len(values) == length:
        return values
    if len(values) == 1:
        return values * length
    return [values[index] if index < len(values) else default for index in range(length)]


def choose_mode(counter: Counter[str], default: str) -> str:
    if not counter:
        return default
    return sorted(counter.items(), key=lambda item: (-item[1], item[0]))[0][0]


def placeholders(count: int) -> str:
    return ", ".join(["%s"] * count)


def create_schema(connection: Connection) -> None:
    """Recreate the unified schema in place; the dataset is rebuilt wholesale.

    Foreign keys stay off during the load so insert order cannot fail on a
    partially built table; schema.find_orphans() re-checks integrity after.
    """
    connection.run(["SET SESSION FOREIGN_KEY_CHECKS = 0"])
    connection.run(schema.drop_statements())
    connection.run(schema.create_statements())
    connection.commit()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    args = parser.parse_args()
    source = args.input.resolve()
    if not source.is_file():
        raise SystemExit(f"找不到源表: {source}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook["简单专利族合并"]
    headers = [text(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    required = set(FAMILY_SOURCE_FIELDS) | set(TECH_SOURCE_FIELDS) | set(OWNER_SOURCE_FIELDS)
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"输入表缺少字段: {missing}")
    position = {name: headers.index(name) for name in required}

    connection = connect()
    create_schema(connection)
    family_sql = f"INSERT INTO families VALUES ({placeholders(3 + len(FAMILY_SOURCE_FIELDS))})"
    tech_sql = f"INSERT INTO family_tech VALUES ({placeholders(14)})"

    families_batch: list[tuple] = []
    tech_batch: list[tuple] = []
    applicant_geo_batch: list[tuple] = []
    comparison_batch: list[tuple] = []
    pending_entities: list[dict] = []
    pending_applicants: list[tuple[str, list[str]]] = []
    family_ids: set[str] = set()
    metrics = Counter()

    # entity_key -> evidence. Standardized owner is used only as a grouping key;
    # the most frequent raw current-owner name remains the display name.
    entity_names: dict[str, Counter[str]] = defaultdict(Counter)
    entity_types: dict[str, Counter[str]] = defaultdict(Counter)
    entity_locations: dict[str, Counter[tuple[str, str, str, str]]] = defaultdict(Counter)
    owner_locations: dict[str, Counter[tuple[str, str, str, str, str]]] = defaultdict(Counter)

    for source_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        family_id = text(row[position["家族ID"]])
        if not family_id or family_id in family_ids:
            raise ValueError(f"家族ID为空或重复，Excel行 {source_row}: {family_id!r}")
        family_ids.add(family_id)
        priority_year = extract_year(
            row[position["最早优先权日"]], row[position["优先权日"]], row[position["首次公开日"]]
        )
        family_values = tuple(text(row[position[field]]) for field in FAMILY_SOURCE_FIELDS)
        families_batch.append((family_id, source_row - 1, priority_year, *family_values))
        technology_code = text(row[position["技术分类代码"]])
        technology_label = text(row[position["技术分类"]])
        if technology_code == "B7" and technology_label == "用聚变支撑技术":
            technology_label = "通用聚变支撑技术"
        tech_batch.append((
            family_id, "complete", text(row[position["技术特征"]]),
            technology_code, technology_label,
            text(row[position["具体技术路线代码"]]), text(row[position["具体技术路线"]]),
            text(row[position["产业链三级代码"]]), text(row[position["产业链一级标签"]]),
            text(row[position["产业链二级标签"]]), text(row[position["产业链三级标签"]]),
            text(row[position["分类依据"]]), "已复核", "",
        ))

        # Applicant geography keeps the applicant-based benchmark convention.
        countries = split_geo_values(row[position["申请人国家/地区"]]) or ["国家未识别"]
        provinces = [
            region for value in split_geo_values(row[position["申请人省市代码"]])
            if (region := clean_region(value)) in CHINA_REGIONS
        ]
        cities = [clean_region(value) for value in split_geo_values(row[position["中国申请人地市"]])]
        districts = [clean_region(value) for value in split_geo_values(row[position["中国申请人区县"]])]
        china_index = 0
        seen_geo: set[tuple[str, str, str, str]] = set()
        for raw_country in countries:
            original, country, implied_province = normalize_country(raw_country)
            province = city = district = ""
            if country == "中国":
                province = implied_province or (provinces[china_index] if china_index < len(provinces) else "")
                city = cities[china_index] if china_index < len(cities) else ""
                district = districts[china_index] if china_index < len(districts) else ""
                if province in {"北京", "天津", "上海", "重庆", "香港", "澳门"} and not city:
                    city = province
                china_index += 1
            key = (country, province, city, district)
            if key in seen_geo:
                continue
            seen_geo.add(key)
            applicant_geo_batch.append((family_id, len(seen_geo), original, country, province, city, district))

        raw_owners = split_values(row[position["当前权利人"]])
        standardized = split_values(row[position["标准化当前权利人"]])
        owner_types = aligned(split_values(row[position["当前权利人企业类型（回填）"]]), len(raw_owners), "待核验")
        owner_countries = aligned(split_values(row[position["当前权利人国家/地区（回填）"]]), len(raw_owners), "国家未识别")
        owner_provinces = aligned(split_values(row[position["当前权利人省/州（仅中国）"]]), len(raw_owners), "不适用")
        owner_cities = aligned(split_values(row[position["当前权利人城市（仅中国）"]]), len(raw_owners), "不适用")
        if raw_owners and len(standardized) != len(raw_owners):
            metrics["标准化当前权利人与当前权利人数量不一致"] += 1
            standardized = raw_owners
        elif not standardized:
            standardized = raw_owners

        entity_source = "当前权利人"
        if not raw_owners:
            fallbacks = (
                ("申请人终属母公司(中文)", row[position["申请人终属母公司(中文)"]]),
                ("申请人终属母公司(英文)", row[position["申请人终属母公司(英文)"]]),
                ("标准化申请人", row[position["标准化申请人"]]),
            )
            entity_source, chosen = next(((name, value) for name, value in fallbacks if text(value)), ("无可用主体字段", "主体未识别"))
            raw_owners = split_values(chosen) or ["主体未识别"]
            standardized = raw_owners
            owner_types = aligned(split_values(row[position["当前权利人企业类型（回填）"]]), len(raw_owners), "主体信息缺失")
            owner_countries = aligned(split_values(row[position["当前权利人国家/地区（回填）"]]), len(raw_owners), "国家未识别")
            owner_provinces = aligned(split_values(row[position["当前权利人省/州（仅中国）"]]), len(raw_owners), "不适用")
            owner_cities = aligned(split_values(row[position["当前权利人城市（仅中国）"]]), len(raw_owners), "不适用")
            metrics["无当前权利人家族"] += 1

        family_entity_items = []
        for index, owner in enumerate(raw_owners):
            canonical = standardized[index] if index < len(standardized) and standardized[index] else owner
            entity_key = normalized(canonical) or normalized(owner) or f"missing:{family_id}:{index}"
            enterprise_type = owner_types[index] if index < len(owner_types) else "待核验"
            if enterprise_type not in VALID_TYPES:
                enterprise_type = "待核验"
            raw_country = owner_countries[index] if index < len(owner_countries) else "国家未识别"
            original_country, country, implied_province = normalize_country(raw_country)
            province = clean_region(owner_provinces[index] if index < len(owner_provinces) else "")
            city = clean_region(owner_cities[index] if index < len(owner_cities) else "")
            if province in {"不适用", "国家未识别"}: province = ""
            if city in {"不适用", "国家未识别"}: city = ""
            if country == "中国":
                province = implied_province or province
                if province in {"北京", "天津", "上海", "重庆", "香港", "澳门"} and not city:
                    city = province
            else:
                province = city = ""
            entity_names[entity_key][owner] += 1
            entity_types[entity_key][enterprise_type] += 1
            entity_locations[entity_key][(original_country, country, province, city)] += 1
            if entity_source == "当前权利人" and owner != "主体未识别":
                owner_locations[owner][(original_country, country, province, city, entity_key)] += 1
            family_entity_items.append((owner, entity_key, entity_source, enterprise_type))
        pending_entities.append({"family_id": family_id, "items": family_entity_items})

        applicants = split_values(row[position["标准化申请人"]]) or split_values(row[position["申请人"]])
        pending_applicants.append((family_id, applicants))
        applicant_set = {normalized(value) for value in applicants if normalized(value)}
        owner_set = {normalized(value) for value in raw_owners if value != "主体未识别" and normalized(value)}
        comparison_batch.append((
            family_id, "; ".join(sorted(applicant_set)), "; ".join(sorted(owner_set)),
            int(bool(applicant_set and owner_set)), int(bool(applicant_set and owner_set and applicant_set == owner_set)),
        ))

        if len(families_batch) >= 1000:
            connection.executemany(family_sql, families_batch)
            connection.executemany(tech_sql, tech_batch)
            connection.executemany("INSERT INTO family_applicant_geographies VALUES (%s,%s,%s,%s,%s,%s,%s)", applicant_geo_batch)
            connection.executemany("INSERT INTO family_entity_comparison VALUES (%s,%s,%s,%s,%s)", comparison_batch)
            families_batch.clear(); tech_batch.clear(); applicant_geo_batch.clear(); comparison_batch.clear()
            # 每个批次落盘，避免整库构建堆积成单个巨型事务。
            connection.commit()

    if families_batch:
        connection.executemany(family_sql, families_batch)
        connection.executemany(tech_sql, tech_batch)
        connection.executemany("INSERT INTO family_applicant_geographies VALUES (%s,%s,%s,%s,%s,%s,%s)", applicant_geo_batch)
        connection.executemany("INSERT INTO family_entity_comparison VALUES (%s,%s,%s,%s,%s)", comparison_batch)

    # Stable entity IDs are assigned by normalized grouping key.
    entity_ids: dict[str, int] = {}
    entity_rows = []
    location_rows = []
    representative_to_key: dict[str, str] = {}
    for entity_id, key in enumerate(sorted(entity_names), start=1):
        representative = choose_mode(entity_names[key], key)
        # Extremely rare standardized collisions must not violate the display-name unique key.
        if representative in representative_to_key and representative_to_key[representative] != key:
            representative = f"{representative}（{entity_id}）"
        representative_to_key[representative] = key
        enterprise_type = choose_mode(entity_types[key], "待核验")
        raw_country, country, province, city = choose_mode(entity_locations[key], ("", "国家未识别", "", ""))
        entity_ids[key] = entity_id
        entity_rows.append((entity_id, representative, enterprise_type, "合并Excel标准化当前权利人归并", None))
        location_rows.append((entity_id, raw_country, country, province, city, "Excel回填字段", sum(entity_locations[key].values())))
    connection.executemany("INSERT INTO entities VALUES (%s,%s,%s,%s,%s)", entity_rows)
    connection.executemany("INSERT INTO entity_locations VALUES (%s,%s,%s,%s,%s,%s,%s)", location_rows)

    family_entity_rows = []
    for item in pending_entities:
        for order, (owner, key, source_name, enterprise_type) in enumerate(item["items"], start=1):
            family_entity_rows.append((
                item["family_id"], order, entity_ids[key], owner, source_name, enterprise_type,
                "Excel字段直接归并" if source_name == "当前权利人" else "Excel回退链",
            ))
    connection.executemany("INSERT INTO family_entities VALUES (%s,%s,%s,%s,%s,%s,%s)", family_entity_rows)

    owner_location_rows = []
    for owner, evidence in sorted(owner_locations.items()):
        raw_country, country, province, city, key = choose_mode(evidence, ("", "国家未识别", "", "", ""))
        owner_location_rows.append((
            owner, entity_ids.get(key), raw_country, country, province, city,
            "Excel回填字段", "已回填", sum(evidence.values()),
        ))
    connection.executemany("INSERT INTO owner_name_locations VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)", owner_location_rows)

    applicant_rows = []
    key_by_normalized_display = {normalized(row[1]): row[0] for row in entity_rows}
    type_by_id = {row[0]: row[2] for row in entity_rows}
    for family_id, applicants in pending_applicants:
        for order, applicant in enumerate(applicants, start=1):
            entity_id = key_by_normalized_display.get(normalized(applicant))
            applicant_rows.append((
                family_id, order, entity_id and applicant or applicant, entity_id,
                type_by_id.get(entity_id, "待核验"), "当前权利人实体同名匹配" if entity_id else "未匹配",
            ))
    connection.executemany("INSERT INTO family_applicants VALUES (%s,%s,%s,%s,%s,%s)", applicant_rows)

    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    metadata = {
        "dataset_title": "IncoPat核聚变统一专利族数据",
        "data_source": "IncoPat consolidated XLSX",
        "family_scope": "简单专利族",
        "relevance_scope": "全部家族纳入；前端按既定规则展示B1-B9",
        "enterprise_grouping": "当前权利人展示；标准化当前权利人仅用于主体归并；缺失时采用回退链",
        "geography_scope": "国家比较采用申请人地理；企业画像采用当前权利人回填地理",
        "historical_country_rule": json.dumps(HISTORICAL_COUNTRY_MAP, ensure_ascii=False, sort_keys=True),
        "generated_at_utc": generated,
        "source_xlsx": str(source),
        "source_xlsx_sha256": sha256(source),
    }
    connection.executemany("INSERT INTO metadata VALUES (%s,%s)", metadata.items())

    quality = [
        ("简单专利族数", str(len(family_ids)), "通过" if len(family_ids) == 68151 else "异常", "家族ID唯一"),
        ("无当前权利人家族数", str(metrics["无当前权利人家族"]), "信息", "按终属母公司/申请人/主体未识别回退"),
        ("标准化权利人数量不一致家族数", str(metrics["标准化当前权利人与当前权利人数量不一致"]), "信息", "不一致时按当前权利人独立归并"),
        ("权利人实体数", str(len(entity_rows)), "信息", "标准化字段后台归并，页面展示当前权利人"),
        ("当前权利人名称数", str(len(owner_location_rows)), "信息", "含地理回填"),
    ]
    connection.executemany("INSERT INTO quality_metrics VALUES (%s,%s,%s,%s)", quality)
    connection.commit()

    # Indexes and views were created up front by create_schema(); re-enable the
    # foreign keys and prove the load left no orphans, replacing sqlite's
    # "PRAGMA integrity_check" gate.
    connection.run(["SET SESSION FOREIGN_KEY_CHECKS = 1"])
    connection.commit()
    orphans = schema.find_orphans(connection)

    checks = {
        "database": config.describe_mysql(),
        "families": connection.execute("SELECT COUNT(*) FROM families").fetchone()[0],
        "tech": connection.execute("SELECT COUNT(*) FROM family_tech").fetchone()[0],
        "entities": connection.execute("SELECT COUNT(*) FROM entities").fetchone()[0],
        "current_owner_families": connection.execute(
            "SELECT COUNT(DISTINCT family_id) FROM family_entities WHERE entity_source='当前权利人'"
        ).fetchone()[0],
        "orphans": orphans,
    }
    connection.close()
    workbook.close()
    if checks["families"] != 68151 or checks["tech"] != checks["families"] or orphans:
        raise RuntimeError(f"构建校验失败: {checks}")
    print(json.dumps({"source": str(source), **checks}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
